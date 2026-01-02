import io, difflib, mimetypes
from typing import Dict, Any, List, Tuple
from PIL import Image
import numpy as np
import fitz  # PyMuPDF
import cv2
from skimage.metrics import structural_similarity as ssim
from docx import Document

# ---------------------------------------------------------------
# File-type detection
# ---------------------------------------------------------------
def detect_filetype(filename: str, data: bytes) -> str:
    mime, _ = mimetypes.guess_type(filename)
    if mime in ["application/pdf"]:
        return "pdf"
    if mime in ["application/vnd.openxmlformats-officedocument.wordprocessingml.document"]:
        return "docx"
    if mime and mime.startswith("image/"):
        return "image"
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data[:2] == b"PK":
        return "docx"
    return "unknown"


# ---------------------------------------------------------------
# PDF utilities
# ---------------------------------------------------------------
def _load_pdf(doc_bytes: bytes) -> fitz.Document:
    return fitz.open(stream=doc_bytes, filetype="pdf")


def _extract_text_spans(doc: "fitz.Document") -> List[Dict[str, Any]]:
    pages = []
    for i, page in enumerate(doc):
        page_dict = page.get_text("dict")
        lines, full_lines = [], []
        for blk in page_dict.get("blocks", []):
            if blk.get("type", 0) != 0:
                continue
            for line in blk.get("lines", []):
                spans = [s for s in line.get("spans", []) if s.get("text", "").strip()]
                if not spans:
                    continue
                text = " ".join(s["text"].strip() for s in spans).strip()
                if not text:
                    continue
                x0 = min(s["bbox"][0] for s in spans)
                y0 = min(s["bbox"][1] for s in spans)
                x1 = max(s["bbox"][2] for s in spans)
                y1 = max(s["bbox"][3] for s in spans)
                lines.append({"text": text, "bbox": [x0, y0, x1, y1]})
                full_lines.append(text)

        pages.append({
            "page_index": i,
            "text": "\n".join(full_lines),
            "lines": lines,
            "w": page.rect.width,
            "h": page.rect.height,
            "_page_obj": page,
        })
    return pages


def _render_page_image(doc: "fitz.Document", page_index: int, dpi: int = 300) -> Image.Image:
    page = doc[page_index]
    zoom = dpi / 72.0
    mat = fitz.Matrix(zoom, zoom)
    pix = page.get_pixmap(matrix=mat, alpha=False)
    return Image.frombytes("L", [pix.width, pix.height], pix.samples)


# ---------------------------------------------------------------
# Text diff and layout drift
# ---------------------------------------------------------------
def _collect_search_boxes(page: "fitz.Page", text: str, max_boxes: int = 12) -> List[List[float]]:
    rects: List[List[float]] = []
    for chunk in [w for w in text.split() if len(w) > 2][:20]:
        try:
            for r in page.search_for(chunk, quads=False):
                rects.append([r.x0, r.y0, r.x1, r.y1])
                if len(rects) >= max_boxes:
                    return rects
        except Exception:
            pass
    return rects


def _text_diff(pageA: Dict[str, Any], pageB: Dict[str, Any]) -> List[Dict[str, Any]]:
    a_lines = pageA["text"].splitlines()
    b_lines = pageB["text"].splitlines()
    sm = difflib.SequenceMatcher(a=a_lines, b=b_lines, autojunk=False)
    diffs = []
    a_boxes = [ln["bbox"] for ln in pageA.get("lines", [])]
    b_boxes = [ln["bbox"] for ln in pageB.get("lines", [])]
    page_obj_A, page_obj_B = pageA.get("_page_obj"), pageB.get("_page_obj")

    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            continue
        before_snip = "\n".join(a_lines[i1:i2])[:400]
        after_snip = "\n".join(b_lines[j1:j2])[:400]
        before_boxes = [a_boxes[k] for k in range(i1, min(i2, len(a_boxes)))]
        after_boxes  = [b_boxes[k] for k in range(j1, min(j2, len(b_boxes)))]
        if not before_boxes and before_snip and page_obj_A:
            before_boxes = _collect_search_boxes(page_obj_A, before_snip)
        if not after_boxes and after_snip and page_obj_B:
            after_boxes = _collect_search_boxes(page_obj_B, after_snip)
        diffs.append({
            "type": "text",
            "page": pageA["page_index"],
            "severity": "Major",
            "description": f"{tag.upper()} on lines A[{i1}:{i2}] vs B[{j1}:{j2}]",
            "meta": {
                "before": before_snip,
                "after": after_snip,
                "before_boxes": before_boxes,
                "after_boxes": after_boxes,
                "page_w": pageA["w"], "page_h": pageA["h"]
            },
        })
    return diffs


def _detect_layout_drift(pageA, pageB, y_threshold: float = 5.0):
    linesA = {ln["text"].strip(): ln["bbox"][1] for ln in pageA.get("lines", []) if ln["text"].strip()}
    linesB = {ln["text"].strip(): ln["bbox"][1] for ln in pageB.get("lines", []) if ln["text"].strip()}
    drift = []
    for txt, yA in linesA.items():
        if txt in linesB and abs(yA - linesB[txt]) > y_threshold:
            bA = next(ln["bbox"] for ln in pageA["lines"] if ln["text"].strip() == txt)
            bB = next(ln["bbox"] for ln in pageB["lines"] if ln["text"].strip() == txt)
            drift.append((bA, bB))
    return drift


# ---------------------------------------------------------------
# Hybrid raster + vector graphic detector
# ---------------------------------------------------------------
def _detect_graphic_differences(docA, docB, page_index, pA, pB, dpi=150):
    """
    Detect genuine new graphics (vector or raster) on a page.
    Conservative for text-heavy PDFs, permissive for sparse layouts.
    """
    pageA, pageB = docA[page_index], docB[page_index]
    scale = dpi / 72.0
    boxes = []

    # ---------- Raster comparison ----------
    pixA = pageA.get_pixmap(dpi=dpi)
    pixB = pageB.get_pixmap(dpi=dpi)
    grayA = np.frombuffer(pixA.samples, np.uint8).reshape(pixA.height, pixA.width, pixA.n)
    grayB = np.frombuffer(pixB.samples, np.uint8).reshape(pixB.height, pixB.width, pixB.n)
    if grayA.shape[2] > 1:
        grayA = cv2.cvtColor(grayA, cv2.COLOR_RGB2GRAY)
        grayB = cv2.cvtColor(grayB, cv2.COLOR_RGB2GRAY)
    score, diff = ssim(grayA, grayB, full=True)
    if score < 0.995:
        diff = (1 - diff) * 255
        diff = diff.astype(np.uint8)
        _, mask = cv2.threshold(diff, 40, 255, cv2.THRESH_BINARY)
        mask = cv2.morphologyEx(mask, cv2.MORPH_CLOSE, np.ones((5,5), np.uint8))
        contours, _ = cv2.findContours(mask, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
        page_area = grayA.shape[0] * grayA.shape[1]
        for c in contours:
            x, y, bw, bh = cv2.boundingRect(c)
            area = bw * bh
            if area < 0.002 * page_area:  # ignore micro changes
                continue
            ratio = max(bw, bh) / (min(bw, bh) + 1e-6)
            if ratio > 25:
                continue
            boxes.append([x / scale, y / scale, (x + bw) / scale, (y + bh) / scale])

    # ---------- Vector comparison ----------
    try:
        vecA, vecB = pageA.get_cdrawings(), pageB.get_cdrawings()
    except Exception:
        vecA, vecB = [], []
    def rects_from_draws(draws):
        rects = []
        for op in draws:
            pts = []
            for seg in op.get("points", []):
                pts.extend(seg)
            if not pts:
                continue
            xs, ys = pts[0::2], pts[1::2]
            r = fitz.Rect(min(xs), min(ys), max(xs), max(ys))
            if r.get_area() > 300:
                rects.append(r)
        return rects
    rA, rB = rects_from_draws(vecA), rects_from_draws(vecB)
    centersA = np.array([[(r.x0 + r.x1)/2, (r.y0 + r.y1)/2] for r in rA]) if rA else np.empty((0,2))
    for r in rB:
        if centersA.size > 0:
            d = np.sqrt(np.sum((centersA - np.array([(r.x0 + r.x1)/2, (r.y0 + r.y1)/2]))**2, axis=1))
            if np.any(d < 15):
                continue
        if r.get_area() > 300 and not any(r.intersects(fitz.Rect(*ln["bbox"])) for ln in pB.get("lines", [])):
            boxes.append([r.x0, r.y0, r.x1, r.y1])

    # deduplicate overlaps
    merged = []
    for b in boxes:
        rect = fitz.Rect(*b)
        if not any(rect.intersects(fitz.Rect(*m)) for m in merged):
            merged.append(b)

    return merged


# ---------------------------------------------------------------
# PDF comparison core
# ---------------------------------------------------------------
def _compare_pdf(before_bytes: bytes, after_bytes: bytes):
    docA = _load_pdf(before_bytes)
    docB = _load_pdf(after_bytes)
    pagesA, pagesB = _extract_text_spans(docA), _extract_text_spans(docB)
    nA, nB, n = len(pagesA), len(pagesB), min(len(pagesA), len(pagesB))
    diffs: List[Dict[str, Any]] = []

    for i in range(n):
        pA, pB = pagesA[i], pagesB[i]
        diffs.extend(_text_diff(pA, pB))
        gboxes = _detect_graphic_differences(docA, docB, i, pA, pB, dpi=300)
        if gboxes:
            diffs.append({
                "type": "graphic",
                "page": i,
                "severity": "Major",
                "description": f"New or modified graphic(s) on page {i+1}",
                "meta": {"graphic_boxes": gboxes, "page_w": pA["w"], "page_h": pA["h"]},
            })
        drift = _detect_layout_drift(pA, pB, y_threshold=5.0)
        for bA, bB in drift:
            diffs.append({
                "type": "layout",
                "page": i,
                "severity": "Minor",
                "description": f"Text block moved vertically on page {i+1}",
                "meta": {"before_boxes": [bA], "after_boxes": [bB], "page_w": pA["w"], "page_h": pA["h"]},
            })

    if nA != nB:
        diffs.append({
            "type": "layout",
            "page": max(0, n - 1),
            "severity": "Major",
            "description": f"Page count changed: {nA} → {nB}",
            "meta": {"pages_before": nA, "pages_after": nB},
        })

    return {
        "pages_before": nA, "pages_after": nB,
        "diffs": diffs,
        "notes": "Hybrid text, layout, raster, and vector graphic difference detection."
    }


# ---------------------------------------------------------------
# DOCX / Image / Unified entry remain unchanged
# ---------------------------------------------------------------
def _compare_docx(before_bytes: bytes, after_bytes: bytes):
    docA, docB = Document(io.BytesIO(before_bytes)), Document(io.BytesIO(after_bytes))
    textA = "\n".join(p.text.strip() for p in docA.paragraphs if p.text.strip())
    textB = "\n".join(p.text.strip() for p in docB.paragraphs if p.text.strip())
    sm = difflib.SequenceMatcher(a=textA.splitlines(), b=textB.splitlines(), autojunk=False)
    diffs = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal": continue
        diffs.append({
            "type": "text", "page": 0, "severity": "Major",
            "description": f"{tag.upper()} lines A[{i1}:{i2}] vs B[{j1}:{j2}]",
            "meta": {"before": "\n".join(textA.splitlines()[i1:i2])[:400],
                     "after": "\n".join(textB.splitlines()[j1:j2])[:400]},
        })
    return {"pages_before": 1, "pages_after": 1, "diffs": diffs, "notes": "DOCX comparison (text-only)."}


def _compare_images(before_bytes: bytes, after_bytes: bytes):
    imgA = Image.open(io.BytesIO(before_bytes)).convert("RGB")
    imgB = Image.open(io.BytesIO(after_bytes)).convert("RGB")
    w, h = min(imgA.width, imgB.width), min(imgA.height, imgB.height)
    imgA, imgB = imgA.resize((w, h)), imgB.resize((w, h))
    grayA, grayB = np.array(imgA.convert("L")), np.array(imgB.convert("L"))
    score, _ = ssim(grayA, grayB, full=True)
    return {
        "pages_before": 1, "pages_after": 1,
        "diffs": [{
            "type": "visual", "page": 0,
            "severity": "Major" if score < 0.995 else "Cosmetic",
            "description": f"Image SSIM={score:.4f}",
            "meta": {"ssim": score, "page_w": w, "page_h": h},
        }],
        "notes": "Image comparison using SSIM.",
    }


async def compare_documents(before_bytes, after_bytes, before_name, after_name):
    ftype = detect_filetype(before_name, before_bytes)
    if ftype == "pdf": return _compare_pdf(before_bytes, after_bytes)
    if ftype == "docx": return _compare_docx(before_bytes, after_bytes)
    if ftype == "image": return _compare_images(before_bytes, after_bytes)
    raise ValueError("Unsupported file type")
